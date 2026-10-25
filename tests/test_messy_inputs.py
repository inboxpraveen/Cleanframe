"""Files as people actually send them: Excel exports, report workbooks, joined exports.

Each test builds its fixture in ``tmp_path`` (no binary files in the repository) and
pins one behaviour found by running such files end to end through the CLI:

* replaying a recipe on the very file it was learned from must not report drift;
* Excel's ``sep=;`` line, JSON Lines in a ``.json``, ``="00123"`` text formulas and
  Excel error values (``#DIV/0!``) are read for what they are;
* header-row hints never point at a data row, and broken workbooks say what they are;
* rows repeating the header are warned about, with a ``skiprows`` that removes them.
"""

from __future__ import annotations

import csv
import warnings
import zipfile

import openpyxl
import pandas as pd
import pytest

import cleanframe as cf
from cleanframe._util import sanitize_csv_value
from cleanframe.cli import _skiprows_arg, main
from cleanframe.codegen import generate_code
from cleanframe.drift import UNPARSED_DATES_KEY, detect_drift
from cleanframe.errors import CleanFrameError
from cleanframe.executor import execute


def _caught(fn, *args, **kwargs):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        out = fn(*args, **kwargs)
    return out, [str(w.message) for w in caught]


def _write_csv(path, rows, *, encoding="utf-8", delimiter=","):
    with open(path, "w", encoding=encoding, newline="") as fh:
        csv.writer(fh, delimiter=delimiter).writerows(rows)
    return path


# --------------------------------------------------------------------- drift on replay
def _dates_file(path, values):
    rows = [["id", "signup"]] + [[f"c{i}", v] for i, v in enumerate(values)]
    return _write_csv(path, rows)


GOOD = ["03/01/2024", "15/02/2024", "28/03/2024", "30/04/2024", "17/05/2024", "21/06/2024"] * 4
NOISE = ["31/02/2024", "45200", "TBD"]  # an invalid date, an Excel serial, a placeholder


def test_replaying_a_recipe_on_its_own_file_reports_no_drift(tmp_path):
    src = _dates_file(tmp_path / "a.csv", GOOD + NOISE)
    result = cf.clean(src, mode="auto")
    learned = result.recipe.source_fingerprint[UNPARSED_DATES_KEY]["signup"]
    assert set(learned) == {"9+/9+/9+", "9+", "A+"}  # shapes only, never the values
    assert not detect_drift(cf.read_frame(src), result.recipe).has_drift
    cf.apply_recipe(src, result.recipe)  # on_drift="error" would raise
    saved = tmp_path / "r.yaml"
    result.recipe.save(saved)
    assert "TBD" not in saved.read_text(encoding="utf-8")
    cf.apply_recipe(src, cf.Recipe.load(saved))


def test_a_new_date_shape_is_still_drift(tmp_path):
    result = cf.clean(_dates_file(tmp_path / "a.csv", GOOD + NOISE), mode="auto")
    nxt = _dates_file(tmp_path / "b.csv", GOOD + ["2024.06.30"])
    report = detect_drift(cf.read_frame(nxt), result.recipe)
    finding = report.by_kind("format_drift")[0]
    assert "new: '2024.06.30'" in finding.message
    with pytest.raises(cf.DriftError):
        cf.apply_recipe(nxt, result.recipe)


def test_a_known_shape_that_becomes_common_is_drift(tmp_path):
    # One stray 31/02 is noise; half the column failing is a DD/MM -> MM/DD switch.
    result = cf.clean(_dates_file(tmp_path / "a.csv", GOOD + ["31/02/2024"]), mode="auto")
    switched = ["01/13/2024", "02/14/2024", "03/15/2024", "04/16/2024", "05/17/2024", "06/18/2024"]
    nxt = _dates_file(tmp_path / "b.csv", GOOD[:6] + switched)
    finding = detect_drift(cf.read_frame(nxt), result.recipe).by_kind("format_drift")[0]
    assert "far more than when the recipe was learned" in finding.message


def test_streamed_replay_agrees_with_whole_frame_replay_on_learned_noise(tmp_path):
    src = _dates_file(tmp_path / "a.csv", GOOD + NOISE)
    result = cf.clean(src, mode="auto", text=True)
    out = tmp_path / "out.csv"
    cf.stream_apply(result.recipe, src, out, chunksize=5)  # raises on drift
    nxt = _dates_file(tmp_path / "b.csv", GOOD + ["2024.06.30"])
    with pytest.raises(cf.DriftError):
        cf.stream_apply(result.recipe, nxt, tmp_path / "out2.csv", chunksize=5)


def test_streaming_reports_parse_losses_once_for_the_whole_file(tmp_path):
    src = _dates_file(tmp_path / "a.csv", GOOD + NOISE * 3)
    recipe = cf.clean(src, mode="auto", text=True).recipe
    summary, notes = _caught(cf.stream_apply, recipe, src, tmp_path / "out.csv", chunksize=4)
    losses = [n for n in notes if "values became missing" in n]
    assert losses == ["CleanFrame: values became missing because an op could not parse them: signup (9)"]
    assert summary.parse_losses == {"signup": 9} and summary.chunks > 5


def test_recipes_without_the_learned_shapes_keep_the_old_behaviour(tmp_path):
    src = _dates_file(tmp_path / "a.csv", GOOD + NOISE)
    recipe = cf.clean(src, mode="auto").recipe
    del recipe.source_fingerprint[UNPARSED_DATES_KEY]
    assert detect_drift(cf.read_frame(src), recipe).by_kind("format_drift")


def test_suggest_update_relearns_the_shapes(tmp_path):
    result = cf.clean(_dates_file(tmp_path / "a.csv", GOOD), mode="auto")
    nxt = _dates_file(tmp_path / "b.csv", GOOD + ["2024.06.30", "TBD"])
    patched, report = cf.suggest_update(nxt, result.recipe)
    assert report.has_drift
    assert not detect_drift(cf.read_frame(nxt), patched).has_drift


# --------------------------------------------------------------------- Excel's sep= line
def test_excel_sep_line_names_the_delimiter(tmp_path):
    p = tmp_path / "eu.csv"
    p.write_text("sep=;\nKunde;Betrag;Datum\nA;1.234,56;01.03.2024\nB;99,10;02.03.2024\n", encoding="utf-8")
    result, notes = _caught(cf.clean, p, mode="auto")
    assert list(result.dataframe.columns) == ["kunde", "betrag", "datum"]
    assert result.dataframe["datum"].tolist() == ["2024-03-01", "2024-03-02"]
    assert result.recipe.read["sep"] == ";" and result.recipe.read["blank_lines"] == 1
    assert any("sep=;" in n for n in notes)
    assert len(cf.apply_recipe(p, result.recipe).dataframe) == 2
    with pytest.raises(CleanFrameError, match=r"sep=;.*header_row=1"):
        cf.read_frame(p)


# --------------------------------------------------------------------- header hints
def test_blank_header_name_error_does_not_point_at_a_data_row(tmp_path):
    p = tmp_path / "h.csv"
    p.write_text("Name,Amount,Amount2,,Notes\nn0,1,2,x,note\nn1,3,4,y,note\n", encoding="utf-8")
    with pytest.raises(CleanFrameError) as err:
        cf.read_frame(p)
    assert "empty column name" in str(err.value) and "header_row" not in str(err.value)


def _report_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "ACME - Q3 sales"
    ws.merge_cells("A1:E1")
    ws.append([])
    ws.append(["Region", "Rep", "Sales", None, "First sale"])  # row 3: upper header
    ws.merge_cells("C3:D3")
    ws.append([None, None, "Units", "Revenue", "Date"])  # row 4: lower header
    for i in range(8):
        ws.append(["North" if i % 3 else None, f"rep{i}", i * 10, i * 2.5, pd.Timestamp("2024-07-01")])
    wb.save(path)
    return path


def test_excel_title_hint_skips_merged_headers_not_data(tmp_path):
    p = _report_workbook(tmp_path / "report.xlsx")
    _, notes = _caught(cf.read_frame, p)
    hint = next(n for n in notes if "header_row=" in n)
    assert "header_row=3" in hint and "span rows 3 and 4" in hint
    fixed = cf.read_frame(p, header_row=3)
    assert list(fixed.columns)[2:4] == ["Units", "Revenue"] and len(fixed) == 8


def test_pivot_output_hint_points_at_the_label_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Sum of Revenue", None, "Quarter", None, None])
    ws.append(["Region", "Product", "Q1", "Q2", "Grand Total"])
    ws.append(["East", "Laptop", 136, 531, 667])
    ws.append([None, "Phone", 401, 331, 732])
    ws.append(["East Total", None, None, None, 1399])
    wb.save(tmp_path / "pivot.xlsx")
    _, notes = _caught(cf.read_frame, tmp_path / "pivot.xlsx")
    assert any("header_row=1 " in n for n in notes)


# --------------------------------------------------------------------- odd formats
def test_json_lines_saved_as_json_is_read(tmp_path):
    p = tmp_path / "rows.json"
    p.write_text('{"a": 1, "b": "x"}\n{"a": 2, "b": "y"}\n', encoding="utf-8")
    df = cf.read_frame(p)
    assert df["a"].tolist() == [1, 2]
    bad = tmp_path / "bad.json"
    bad.write_text('{"a": 1,', encoding="utf-8")
    with pytest.raises(CleanFrameError):
        cf.read_frame(bad)


@pytest.mark.parametrize(
    ("name", "content", "expected"),
    [
        ("export.xls", b"<html><table><tr><td>a</td></tr></table></html>", "HTML or XML table"),
        ("export.xls", b'<?xml version="1.0"?><Workbook/>', "HTML or XML table"),
        ("locked.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600, "password-protected"),
    ],
)
def test_files_that_are_not_workbooks_say_what_they_are(tmp_path, name, content, expected):
    p = tmp_path / name
    p.write_bytes(content)
    with pytest.raises(CleanFrameError, match=expected):
        cf.read_frame(p)


def test_zip_that_is_not_a_workbook(tmp_path):
    p = tmp_path / "x.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(CleanFrameError, match="zip archive without the parts Excel writes"):
        cf.read_frame(p)


# --------------------------------------------------------------------- joined exports
def test_rows_repeating_the_header_are_warned_with_a_working_skiprows(tmp_path):
    header = ["City", "Amount"]
    rows = [header] + [["Pune", "10"]] * 3 + [header] + [["Delhi", "20"]] * 2
    p = _write_csv(tmp_path / "joined.csv", rows)
    _, notes = _caught(cf.read_frame, p)
    msg = next(n for n in notes if "repeat the header" in n)
    assert "skiprows=[4]" in msg and "--skiprows 4," in msg
    assert len(cf.read_frame(p, skiprows=[4])) == 5
    assert _skiprows_arg("4,") == [4] and _skiprows_arg("4") == 4 and _skiprows_arg("4,9") == [4, 9]
    out = tmp_path / "out.csv"
    assert main(["clean", str(p), "--skiprows", "4,", "--out", str(out), "--mode", "auto"]) == 0
    assert set(pd.read_csv(out)["city"]) == {"Pune", "Delhi"}  # no "City" row cleaned as data


def test_dedup_on_a_key_never_merges_rows_that_have_no_key(tmp_path):
    # Two LLMs proposed dedup on "email" for a customer export where a third of the
    # emails were blank; pandas counts missing keys as equal, so 55 different
    # customers were dropped as "duplicates" of the first one.
    df = pd.DataFrame({
        "name": ["a", "b", "c", "d", "e", "f"],
        "email": ["x@y.com", None, "X@Y.com ", "  ", float("nan"), "x@y.com"],
    })
    recipe = cf.Recipe.from_dict({"version": 1, "columns": {},
                                  "dedup": {"subset": ["email"], "ignore_case": True}})
    out = execute(recipe, df).dataframe
    assert out["name"].tolist() == ["a", "b", "d", "e"]  # c and f repeat a's email
    ns: dict = {}
    exec(compile(generate_code(recipe), "g.py", "exec"), ns)
    assert ns["clean"](df)["name"].tolist() == ["a", "b", "d", "e"]
    full_row = cf.Recipe.from_dict({"version": 1, "columns": {}, "dedup": {}})
    blank_twins = pd.DataFrame({"a": [None, None], "b": [None, None]})
    assert len(execute(full_row, blank_twins).dataframe) == 1  # identical rows still dedupe


def test_normal_files_get_no_repeated_header_warning(tmp_path):
    p = _write_csv(tmp_path / "ok.csv", [["City", "Amount"], ["Pune", "10"], ["Delhi", "20"]])
    _, notes = _caught(cf.clean, p)
    assert not any("repeat the header" in n for n in notes)


# --------------------------------------------------------------------- Excel cell idioms
def test_excel_error_values_become_missing_with_their_own_message(tmp_path):
    df = pd.DataFrame({
        "item": ["a", "b", "c", "d"],
        "price": ["10.5", "#REF!", "7.25", "#VALUE!"],
        "lookup": ["Found", "#N/A", "#NAME?", "Found"],
    })
    result = cf.clean(df)
    messages = [i.message for i in result.issues if i.column == "price" and i.kind == "disguised_nulls"]
    assert messages and "Excel error values" in messages[0] and "failed" in messages[0]
    assert result.dataframe["price"].isna().tolist() == [False, True, False, True]
    assert result.dataframe["lookup"].isna().tolist() == [False, True, True, False]


def test_excel_text_formulas_are_unwrapped_and_codegen_agrees(tmp_path):
    df = pd.DataFrame({"code": ['="00123"', '="04500"', '="00007"'], "n": ["a", "b", "c"]})
    result = cf.clean(df, mode="auto")
    assert result.dataframe["code"].tolist() == ["00123", "04500", "00007"]
    ns: dict = {}
    exec(compile(generate_code(result.recipe), "g.py", "exec"), ns)
    pd.testing.assert_frame_equal(
        ns["clean"](df).reset_index(drop=True),
        execute(result.recipe, df, mode="auto").dataframe.reset_index(drop=True),
    )
    mostly_plain = pd.DataFrame({"f": ['="x"', "plain", "text", "more"]})
    assert not [i for i in cf.clean(mostly_plain).issues if i.kind == "excel_text_formula"]


@pytest.mark.parametrize("value", ["-3.5%", "+5%", "-1,234.50", "-1e5", "+919876543210"])
def test_signed_numbers_are_not_escaped_as_formulas(value):
    assert sanitize_csv_value(value) == value


@pytest.mark.parametrize("value", ["=1+1", "-2+3", "@SUM(1)", "+cmd|' /C calc'!A0", "-1,23", "--"])
def test_formula_like_values_are_still_escaped(value):
    assert sanitize_csv_value(value) == "'" + value
